"""Extract text from supported syllabus file formats."""
from pathlib import Path
from io import BytesIO
import csv
import pymupdf as fitz
import re
from datetime import date
from docx import Document
from docx.table import Table as WordTable
from fastapi import HTTPException


class ParsedDocument(str):
    """Keep geometric table rows alongside the original text for local extraction."""
    def __new__(cls, text, schedule_rows):
        obj = super().__new__(cls, text)
        obj.schedule_rows = schedule_rows
        return obj


def _schedule_rows(doc, text):
    # A year must be explicit; never guess the academic year from the computer clock.
    year_match = re.search(r"\b(?:fall|spring|summer|winter)\s+((?:19|20)\d{2})\b", text, re.I) or re.search(r"\b(?:19|20)\d{2}\b", text[:2000])
    if not year_match:
        return []
    year = int(year_match.group(1) if year_match.lastindex else year_match.group())
    course = next((x.strip() for x in text.splitlines() if re.search(r"\b[A-Z]{2,}\s*\d{3,}", x)), "Unknown course")
    months = {name: i for i, name in enumerate(['jan','feb','mar','apr','may','jun','jul','aug','sep','oct','nov','dec'], 1)}
    clean = lambda value: re.sub(r"\s+", " ", value or '').strip()
    result, seen = [], set()
    for page in doc:
        try:
            tables = [page] if isinstance(page, WordTable) else page.find_tables().tables
        except Exception:
            continue  # Text-only parsing still handles non-tabular documents.
        for table in tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows] if isinstance(table, WordTable) else table.extract()
            header_index = next((i for i, row in enumerate(rows) if any('due date' in clean(c).lower() for c in row)), None)
            if header_index is None:
                continue
            header = rows[header_index]
            due_col = next(i for i,c in enumerate(header) if 'due date' in clean(c).lower())
            date_col = next((i for i,c in enumerate(header) if clean(c).lower() == 'date'), None)
            if date_col is None:
                continue
            exam_col = next((i for i,c in enumerate(header) if 'exam' in clean(c).lower()), len(header))
            task_cols = [(i, clean(c)) for i,c in enumerate(header) if re.search(r'attendance|atetndance|practices|scenarios|programming|discussion|exam', clean(c), re.I)]
            # PDF geometry sometimes splits AM/PM and week numbers into a continuation row.
            merged = []
            for row in rows[header_index+1:]:
                if merged and not clean(row[date_col]) and not any(clean(row[i]) for i,_ in task_cols):
                    merged[-1] = [' '.join(filter(None,[a,b])) for a,b in zip(merged[-1],row)]
                else:
                    merged.append(row)
            for row in merged:
                due_text = clean(' '.join(c or '' for c in row[due_col:exam_col]))
                numeric = re.search(r'\b(\d{1,2})[/-](\d{1,2})[/-](\d{4}|\d{2})\b', due_text)
                match = re.search(r'\b([A-Za-z]+)\.?\s+(\d{1,2})\b', due_text)
                try:
                    if numeric:
                        explicit_year = int(numeric[3])
                        if explicit_year < 100:
                            explicit_year += (year // 100) * 100
                        due = date(explicit_year,int(numeric[1]),int(numeric[2])).isoformat()
                    elif match and match[1].lower()[:3] in months:
                        due = date(year,months[match[1].lower()[:3]],int(match[2])).isoformat()
                    else:
                        continue
                except ValueError:
                    continue
                reading = next((clean(row[i]) for i,c in enumerate(header) if 'reading' in clean(c).lower()), '')
                chapter = re.search(r'Chapter\s+\d+', reading, re.I)
                exam_cells = [(i, clean(c)) for i,c in enumerate(row) if i > due_col and re.match(r'^Exam\s+(?:[IVX]+|\d+)\b', clean(c), re.I)]
                exam = exam_cells[-1][1] if exam_cells else ''
                row_task_cols = task_cols + [(i, 'Exams') for i,_ in exam_cells if i not in [col for col,_ in task_cols]]
                time = re.search(r'(\d{1,2}:\d{2}\s*(?:AM|PM))', due_text, re.I)
                for col, label in row_task_cols:
                    title = clean(row[col])
                    if not title or re.search(r'no assignments|college closed|labor day|thanksgiving|schedule to be announced|project sheet assigned|overview of project', title, re.I):
                        continue
                    if 'exam' in label.lower() and not re.search(r'\bexam\b', title, re.I):
                        continue
                    if 'attendance' in label.lower() or 'atetndance' in label.lower():
                        if exam and re.fullmatch(r'Exam\s+[IVX\d]+', title, re.I):
                            continue
                        if exam:
                            title = re.sub(r'\s*&\s*Exam\s+[IVX\d]+', '', title, flags=re.I)
                    if chapter and re.match(r'Practices?\b', title, re.I):
                        title = chapter.group() + ' - ' + title
                    titles = [x.strip() for x in re.split(r'(?=Chapter\s+\d+\s+Question\s+\d+)', title, flags=re.I) if x.strip()]
                    for title in titles:
                        key = (title.casefold(), due)
                        if key in seen:
                            continue
                        seen.add(key)
                        result.append(dict(courseName=course[:200],title=title[:300],dueDate=due,
                            instructions=('Due at ' + time[1] + '. Timezone is not specified in the document. ' if time else '') + 'Imported from the weekly schedule; verify this item and its deadline.',
                            sourceQuote=(label + ': ' + clean(row[col]) + ' | Due: ' + due_text)[:500]))
    return result



def _weekly_text_rows(text, filename):
    """Read weekly schedules whose borderless PDF tables cannot be reconstructed."""
    month_pattern = r'(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)'
    deadline = re.compile(r'\b(' + month_pattern + r')\.?\s+(\d{1,2}),?\s+((?:19|20)\d{2}),?\s*(?:(?:at|@)\s*)?(\d{1,2}:\d{2}\s*(?:am|pm))?', re.I)
    months={name:i for i,name in enumerate(['jan','feb','mar','apr','may','jun','jul','aug','sep','oct','nov','dec'],1)}
    course_match=re.search(r'\b([A-Z]{2,})[ _-]*(\d{3})\b', Path(filename).stem)
    course=' '.join(course_match.groups()) if course_match else 'Unknown course'
    result=[]; seen=set()
    for section in re.split(r'(?im)^\s*Week\s+\d+\s*$',text)[1:]:
        # Undated final-exam notices are not deadlines for the preceding week's work.
        section=re.split(r'(?im)^.*FINAL EXAMS SCHEDULE TO BE ANNOUNCED.*$',section)[0]
        match=deadline.search(section)
        if not match: continue
        try: due=date(int(match[3]),months[match[1][:3].lower()],int(match[2])).isoformat()
        except ValueError: continue
        body=deadline.sub('',section)
        titles=[]
        for line in body.splitlines():
            line=re.sub(r'\s+',' ',line).strip()
            if not line or re.fullmatch(r'Chapter\s+\d+',line,re.I): continue
            if re.match(r'^'+month_pattern+r'\.?\s+\d',line,re.I): continue
            if re.search(r'break|no class|final exam|schedule to be announced',line,re.I): continue
            if re.match(r'^(?:Attendance Verification Quiz|Student ID|Practice Test|Discussion\s+\d|Lab Sim|Virtual Lab|Chapter\s+\d+\s+Scenario|Knowledge Check|Pre-Assessment|Can You)',line,re.I):
                titles.append(line)
            elif titles:
                titles[-1]+=' '+line
        for title in titles:
            key=(title.casefold(),due)
            if key in seen: continue
            seen.add(key)
            result.append(dict(courseName=course,title=title[:300],dueDate=due,
                instructions=('Due at '+match[4]+'. ' if match[4] else '')+'From a weekly schedule; verify the deadline and requirements. Timezone not specified.',
                sourceQuote=(title+' | Due: '+re.sub(r'\s+',' ',match[0]))[:500]))
    return result


SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv"}

def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise HTTPException(status_code=415, detail="Unsupported file type. Upload PDF, DOCX, TXT, MD, or CSV.")
    schedule_rows = []
    try:
        if ext == ".pdf":
            with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                text = "\n".join(page.get_text("text") for page in doc)
                schedule_rows = _schedule_rows(doc, text) or _weekly_text_rows(text, filename)
        elif ext == ".docx":
            doc = Document(BytesIO(file_bytes))
            chunks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(chunks)
            schedule_rows = _schedule_rows(doc.tables, text + "\n" + Path(filename).stem.replace("_", " "))
        elif ext == ".csv":
            decoded = file_bytes.decode("utf-8-sig")
            text = "\n".join(" | ".join(row) for row in csv.reader(decoded.splitlines()))
        else:
            text = file_bytes.decode("utf-8-sig")
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read this document: {exc}") from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise HTTPException(status_code=422, detail="No readable text was found. This may be a scanned PDF; OCR is not included yet.")
    return ParsedDocument(text, schedule_rows)
