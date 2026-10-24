from services.document_parser import _schedule_rows, ParsedDocument
from services.ai_service import _fallback_extract

class Table:
    def extract(self):
        return [
            ['Week','Date','Attendance','Reading','Practices & Quiz','Scenarios','Due Dates','Exams &'],
            ['Week 1','Aug. 24 - Aug. 30','Chapter 1 Review Questions','Chapter 1','Practices: 1, 2 & 3','Grandfield College Page 18','Aug. 30 @ 11:59',''],
            [None,None,None,None,None,None,'PM',None],
            ['Week 12','Nov. 9 - Nov. 15','Exam II','','Project Sheet Assigned Overview of Project','','Nov 15 @ 11:59 PM','Exam II - Chapters 5 - 8'],
            ['','','','','','','','College Closed'],
        ]
class Page:
    def find_tables(self):
        return type('Tables',(),{'tables':[Table()]})()

def test_table_columns_keep_deadlines_and_exclude_non_assignments():
    rows=_schedule_rows([Page()], 'CPT242 Fall 2026 Schedule - Online')
    items=_fallback_extract(ParsedDocument('Original PDF text',rows),None)
    assert len(items)==4
    assert items[0].dueDate=='2026-08-30'
    assert '11:59 PM' in items[0].instructions
    assert items[-1].title=='Exam II - Chapters 5 - 8'
    assert items[-1].dueDate=='2026-11-15'
    assert all(x.needsReview for x in items)
    assert not any('Project Sheet Assigned' in x.title for x in items)

def test_table_does_not_guess_year():
    assert _schedule_rows([Page()], 'CPT242 Schedule')==[]


def test_word_schedule_numeric_dates_and_multiple_questions():
    from io import BytesIO
    from docx import Document
    from services.document_parser import extract_text
    from services.ai_service import extract_assignments
    doc=Document()
    doc.add_paragraph('IST 271 Fall 2026 Schedule')
    table=doc.add_table(rows=1,cols=7)
    for cell,value in zip(table.rows[0].cells,['Week','Date','Reading','Programming Problems','Discussions','Due Dates','Exams']): cell.text=value
    for values in [
        ['Week 1','Aug 24 - 30','Chapter 7','','Discussion 1','9-2-26 at 11:59PM',''],
        ['Week 13','Nov 16 - 22','','Chapter 11 Question 1\nChapter 11 Question 2','Discussion 13','12-2-26 at 11:59PM',''],
        ['Week 14','Nov 23 - 29','','Thanksgiving Week','','12-2-26 at 11:59PM',''],
        ['Finals','Dec 8 - 11','','','','12-9-2026','Exam 3'],
    ]:
        for cell,value in zip(table.add_row().cells,values): cell.text=value
    buffer=BytesIO();doc.save(buffer)
    text=extract_text(buffer.getvalue(),'schedule.docx')
    items=_fallback_extract(text,None)
    assert len(items)==5
    assert items[0].dueDate=='2026-09-02'
    assert any(x.title=='Chapter 11 Question 2' and x.dueDate=='2026-12-02' for x in items)
    assert items[-1].title=='Exam 3'
    assert items[-1].dueDate=='2026-12-09'
    assert not any('Thanksgiving' in x.title for x in items)


def test_pdf_rows_are_not_word_cells():
    class PdfTable:
        rows=[[(0,0,1,1)]]
        def extract(self): return [['FALL 2026 SCHEDULE']]
    class PdfPage:
        def find_tables(self): return type('Tables',(),{'tables':[PdfTable()]})()
    assert _schedule_rows([PdfPage()], 'FALL 2026 SCHEDULE')==[]


def test_borderless_weekly_pdf_keeps_deadline_and_wrapped_title():
    from services.document_parser import _weekly_text_rows
    text='''FALL 2026 SCHEDULE
Week 8
Oct 12 - 18
Chapter 8
Practice Test: Chapter 8
Discussion 8
Virtual Lab 8-1
Oct 25, 2026, at
11:59 pm
Chapter 8 Scenario
Week 15
Nov 30 - Dec 6
Can You Create, Manage, and
Dec 6, 2026, at 11:59 pm
Monitor Azure VM and Storage
Dec 8 - 11
FINAL EXAMS SCHEDULE TO BE ANNOUNCED
Final Exam
'''
    rows=_weekly_text_rows(text,'IST 196 Schedule.pdf')
    assert len(rows)==5
    assert rows[0]['dueDate']=='2026-10-25'
    assert rows[-1]['title']=='Can You Create, Manage, and Monitor Azure VM and Storage'
    assert rows[-1]['dueDate']=='2026-12-06'
    assert all(row['courseName']=='IST 196' for row in rows)
